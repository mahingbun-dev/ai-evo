"""
将专利补充内容（技术实现细节）Markdown 转换为 docx 文件
统一格式：正文宋体10.5pt，代码块Courier New 9pt灰色，标题统一字号加粗
"""
import re
from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


# ── 全局格式常量 ──
FONT_NORMAL = '宋体'
FONT_CODE = 'Courier New'
SIZE_H1 = Pt(16)
SIZE_H2 = Pt(14)
SIZE_H3 = Pt(12)
SIZE_H4 = Pt(11)
SIZE_NORMAL = Pt(10.5)
SIZE_CODE = Pt(9)
SIZE_TABLE = Pt(10)
COLOR_CODE = RGBColor(51, 51, 51)


def set_run_format(run, size=SIZE_NORMAL, bold=False, font_name=FONT_NORMAL, color=None):
    """统一设置 run 格式"""
    run.font.size = size
    run.font.name = font_name
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_heading_paragraph(doc, text, level=1):
    """添加标题段落"""
    p = doc.add_paragraph()
    run = p.add_run(text)
    sizes = {1: SIZE_H1, 2: SIZE_H2, 3: SIZE_H3, 4: SIZE_H4}
    set_run_format(run, size=sizes.get(level, SIZE_H4), bold=True)
    return p


def add_rich_paragraph(doc, text):
    """添加含加粗标记的正文段落，统一字体和字号"""
    p = doc.add_paragraph()
    # 按 **...** 分割，奇数索引为加粗
    parts = re.split(r'\*\*(.*?)\*\*', text)
    for idx, part in enumerate(parts):
        if not part:
            continue
        run = p.add_run(part)
        is_bold = (idx % 2 == 1)
        set_run_format(run, size=SIZE_NORMAL, bold=is_bold)
    return p


def add_code_paragraphs(doc, code_text):
    """添加代码块：每行一个段落，统一 Courier New 9pt 灰色"""
    lines = code_text.rstrip('\n').split('\n')
    for line in lines:
        p = doc.add_paragraph()
        run = p.add_run(line if line else ' ')
        set_run_format(run, size=SIZE_CODE, font_name=FONT_CODE, color=COLOR_CODE)


def parse_markdown_to_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()
    style = doc.styles['Normal']
    style.font.size = SIZE_NORMAL
    style.font.name = FONT_NORMAL

    i = 0
    code_block = None
    in_table = False
    table_rows = []

    def flush_table():
        nonlocal in_table, table_rows
        if not in_table or not table_rows:
            in_table = False
            table_rows = []
            return
        # 统一去掉 markdown 表格加粗标记
        cleaned = []
        for row in table_rows:
            cleaned.append([re.sub(r'\*\*(.*?)\*\*', r'\1', c) for c in row])

        cols = max(len(r) for r in cleaned)
        table = doc.add_table(rows=len(cleaned), cols=cols)
        table.style = 'Table Grid'
        for ri, row_data in enumerate(cleaned):
            for ci, cell_text in enumerate(row_data):
                if ci < cols:
                    cell = table.cell(ri, ci)
                    cell.text = ''
                    run = cell.paragraphs[0].runs[0] if cell.paragraphs[0].runs else cell.paragraphs[0].add_run(cell_text)
                    if not cell.paragraphs[0].runs:
                        pass
                    else:
                        cell.paragraphs[0].runs[0].text = cell_text
                    # 重新设置 cell 内文字格式
                    cell.text = cell_text
                    for paragraph in cell.paragraphs:
                        for r in paragraph.runs:
                            set_run_format(r, size=SIZE_TABLE)
        in_table = False
        table_rows = []
        doc.add_paragraph()

    while i < len(lines):
        line = lines[i]
        stripped = line.rstrip('\n')

        # ── 代码块 ──
        if stripped.startswith('```'):
            if code_block is not None:
                add_code_paragraphs(doc, code_block)
                code_block = None
            else:
                flush_table()
                code_block = ''
            i += 1
            continue

        if code_block is not None:
            code_block += stripped + '\n'
            i += 1
            continue

        # ── 表格 ──
        if '|' in stripped and stripped.startswith('|'):
            if re.match(r'^\|[\s\-:|]+\|$', stripped):
                i += 1
                continue
            cells = [c.strip() for c in stripped.split('|')[1:-1]]
            if not in_table:
                in_table = True
                table_rows = []
            table_rows.append(cells)
            i += 1
            continue
        else:
            if in_table:
                flush_table()

        # ── 标题 (# ~ ####) ──
        heading_match = re.match(r'^(#{1,4})\s+(.*)', stripped)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2).strip()
            add_heading_paragraph(doc, text, level=level)
            i += 1
            continue

        # ── 分隔线 ──
        if stripped.strip() in ('---', '***', '___'):
            doc.add_paragraph()
            i += 1
            continue

        # ── 空行 ──
        if not stripped.strip():
            i += 1
            continue

        # ── 普通文本（可能含 **加粗**） ──
        add_rich_paragraph(doc, stripped)
        i += 1

    flush_table()

    doc.save(docx_path)
    print(f"文档已成功生成: {docx_path}")

    # 验证
    verify_doc = Document(docx_path)
    total = len(verify_doc.paragraphs)
    non_empty = sum(1 for p in verify_doc.paragraphs if p.text.strip())
    print(f"总段落数: {total}, 非空段落数: {non_empty}")

    # 检查格式一致性
    font_sizes = set()
    font_names = set()
    for p in verify_doc.paragraphs:
        if not p.text.strip():
            continue
        for r in p.runs:
            if r.font.size:
                font_sizes.add(r.font.size)
            if r.font.name:
                font_names.add(r.font.name)
    print(f"使用的字号种类: {len(font_sizes)} — {sorted(font_sizes)}")
    print(f"使用的字体种类: {len(font_names)} — {sorted(font_names)}")

    full_text = "\n".join(p.text for p in verify_doc.paragraphs)
    checks = [
        "多引擎协同评测架构", "变量注入模块", "系统内部数据流转",
        "模型调用结构", "评测指令动态模板", "数据集快照机制",
        "异步任务调度", "各模块输入输出", "创造性与实用性",
        "VariableContext", "JudgeResponseParser", "ModelCallerSync", "双重中断机制",
    ]
    for c in checks:
        found = c in full_text
        print(f"  {'OK' if found else 'MISSING'}: {c}")


if __name__ == "__main__":
    md_path = "docs/专利补充内容_技术实现细节.md"
    docx_path = "docs/专利补充内容_技术实现细节.docx"
    parse_markdown_to_docx(md_path, docx_path)
